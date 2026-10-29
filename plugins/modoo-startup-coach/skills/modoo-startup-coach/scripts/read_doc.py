#!/usr/bin/env python3
"""신청서·사업계획서·발표자료 파일에서 글을 꺼낸다 (PDF, DOCX, PPTX, HWPX, TXT).

사용법:
  python3 read_doc.py 파일 [--images 폴더] [--group 3]
    --images : PDF 쪽을 PNG로 저장 (그림 PDF일 때)
    --group  : 쪽 이미지를 N장씩 세로로 붙인 group_01.png… 도 만든다 (발표자료 훑어보기)
"""
import argparse
import os
import re
import sys
import zipfile

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")


def group_images(image_dir, n):
    """쪽 이미지를 n장씩 세로로 붙여 group_01.png… 로 저장한다 (발표자료를 빨리 훑을 때)."""
    try:
        from PIL import Image
    except ImportError:
        print("(묶음 이미지는 Pillow가 있어야 해요: pip install pillow)", file=sys.stderr)
        return []
    files = sorted(f for f in os.listdir(image_dir) if f.startswith("page_") and f.endswith(".png"))
    out = []
    for k in range(0, len(files), n):
        ims = [Image.open(os.path.join(image_dir, f)) for f in files[k:k + n]]
        w = max(im.width for im in ims)
        sheet = Image.new("RGB", (w, sum(im.height for im in ims)), "white")
        y = 0
        for im in ims:
            sheet.paste(im, (0, y))
            y += im.height
        name = os.path.join(image_dir, f"group_{k // n + 1:02d}.png")
        sheet.save(name)
        out.append(name)
    return out


def read_pdf(path, image_dir=None):
    out, pages = [], 0
    try:
        try:
            import pymupdf as fitz
        except ImportError:
            import fitz
        d = fitz.open(path)
        pages = len(d)
        for i, p in enumerate(d, 1):
            out.append(f"\n=== {i}쪽 ===\n{p.get_text()}")
            if image_dir:
                os.makedirs(image_dir, exist_ok=True)
                p.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(os.path.join(image_dir, f"page_{i:02d}.png"))
    except ImportError:
        try:
            from pypdf import PdfReader
        except ImportError:
            sys.exit("PDF를 읽을 라이브러리가 없어요: pip install pymupdf 또는 pip install pypdf")
        r = PdfReader(path)
        pages = len(r.pages)
        for i, p in enumerate(r.pages, 1):
            out.append(f"\n=== {i}쪽 ===\n{p.extract_text() or ''}")
        if image_dir:
            print("(쪽 이미지는 pymupdf가 있어야 만들 수 있어요: pip install pymupdf)", file=sys.stderr)
    text = "".join(out)
    chars = len(re.sub(r"\s|=== \d+쪽 ===", "", text))
    if pages and chars / pages < 80:
        text += (f"\n\n[안내] 글이 거의 없어요(쪽당 약 {chars // max(pages, 1)}자). 그림으로 된 PDF 같아요. "
                 "--images 폴더 옵션으로 쪽 이미지를 만들어 직접 보세요.")
    return text


def read_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(path)
    out = []
    for child in d.element.body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            t = Paragraph(child, d).text
            if t.strip():
                out.append(t)
        elif tag == "tbl":
            tbl = Table(child, d)
            out.append("[표]")
            for row in tbl.rows:
                cells = []
                prev = None
                for c in row.cells:
                    if c._tc is prev:
                        continue
                    prev = c._tc
                    cells.append(c.text.replace("\n", " / ").strip())
                out.append(" | ".join(cells))
            out.append("[/표]")
    return "\n".join(out)


def read_pptx(path):
    try:
        from pptx import Presentation
        prs = Presentation(path)
        out = []
        for i, slide in enumerate(prs.slides, 1):
            out.append(f"\n=== {i}장 ===")
            for shp in slide.shapes:
                if shp.has_text_frame and shp.text_frame.text.strip():
                    out.append(shp.text_frame.text)
                if getattr(shp, "has_table", False) and shp.has_table:
                    for row in shp.table.rows:
                        out.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                out.append("[발표자 노트] " + slide.notes_slide.notes_text_frame.text)
        return "\n".join(out)
    except ImportError:
        z = zipfile.ZipFile(path)
        names = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                       key=lambda n: int(re.findall(r"\d+", n)[-1]))
        out = []
        for i, n in enumerate(names, 1):
            xml = z.read(n).decode("utf-8", "ignore")
            out.append(f"\n=== {i}장 ===\n" + " ".join(re.findall(r"<a:t>([^<]*)</a:t>", xml)))
        return "\n".join(out)


def read_hwpx(path):
    z = zipfile.ZipFile(path)
    names = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml$", n))
    out = []
    for n in names:
        xml = z.read(n).decode("utf-8", "ignore")
        for para in re.findall(r"<hp:p[ >].*?</hp:p>", xml, flags=re.S):
            t = "".join(re.findall(r"<hp:t[^>]*>([^<]*)</hp:t>", para))
            if t.strip():
                out.append(t)
    return "\n".join(out) or "[안내] HWPX에서 글을 찾지 못했어요. PDF나 DOCX로 저장해 다시 올려 주세요."


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("file")
    ap.add_argument("--images", help="PDF 쪽 이미지 저장 폴더")
    ap.add_argument("--group", type=int, default=0, help="쪽 이미지를 N장씩 묶기")
    a = ap.parse_args()
    ext = os.path.splitext(a.file)[1].lower()
    if ext == ".pdf":
        print(read_pdf(a.file, a.images))
        if a.images and a.group > 0:
            made = group_images(a.images, a.group)
            if made:
                print(f"\n[묶음 이미지 {len(made)}장] " + ", ".join(os.path.basename(m) for m in made))
    elif ext == ".docx":
        print(read_docx(a.file))
    elif ext == ".pptx":
        print(read_pptx(a.file))
    elif ext == ".hwpx":
        print(read_hwpx(a.file))
    elif ext == ".hwp":
        print("[안내] HWP(옛 한글 형식)는 바로 읽기 어려워요. 한글에서 PDF나 DOCX로 저장해 다시 올려 주세요.")
    else:
        with open(a.file, encoding="utf-8-sig", errors="replace") as f:
            print(f.read())


if __name__ == "__main__":
    main()
